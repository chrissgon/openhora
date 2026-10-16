#!/usr/bin/env python3
"""The program that runs INSIDE the converter's container (providers/converter/Dockerfile copies it to /opt/converter): it
converts one document to text with adopted open tools and prints the result on its standard output as frames. It is not
imported by the host except in the tests, and it never runs on the host in production.

What it uses, all of it adopted and none of it written here: poppler's `pdftotext` for a PDF, pandoc for a Word document and
an HTML page, openpyxl for a spreadsheet, python-pptx for a presentation (a small function here writes its slides as Markdown), and the
standard library for CSV and plain text. Each tool reads the one input file;
the container has no network, a read-only root file system, no capability, and a memory-backed /scratch, and the host copies
back only what this program prints.

  python3 -I _inner.py --type <pdf|docx|xlsx|pptx|ooxml|html|csv|text> --input <file> [--work <folder>]

The frames (standard output only, nothing else is ever printed there):

  OHF1\\n
  FILE <name> <size>\\n<size bytes>          one per output: converted.md or converted.csv, and info.json
  END\\n

The host (providers/converter/container.py) parses them with fixed rules: the names, the sizes and the count are checked
there, not trusted from here. A conversion that fails prints nothing on standard output, a reason on standard error and exits
with a code that is not 0.

A Word document, a spreadsheet or a presentation that carries macros (a VBA project part, a macro sheet, a macro-enabled content type) is
refused whole, before any tool reads it, with one fixed sentence (MACRO_REFUSAL) and exit 3, whatever the file is called.

The converted text is cut at CAP + 1 bytes, so that the host sees it is over the cap and refuses the file; the host's cap is the
one that counts. Standard library, plus openpyxl for a spreadsheet and python-pptx for a presentation (both in the image's virtual environment).
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import secrets
import shutil
import subprocess
import sys
import zipfile

CAP = 1024 * 1024                 # providers/converter/container.py CONVERTED_CAP and runtime/sources.py CONVERTED_CAP
INFO_CAP = 16 * 1024
TOOL_SECONDS = 120
ZIP_ENTRIES_MAX = 5000
ZIP_BYTES_MAX = 300 * 1024 * 1024     # the sum of the uncompressed sizes a package may declare: a zip bomb is refused before any tool reads it
PPTX_BYTES_MAX = 150 * 1024 * 1024     # a presentation: python-pptx keeps every part in memory (the container has 1024 MiB), so its caps are lower
PPTX_XML_MAX = 24 * 1024 * 1024        # and the XML a library parses into a tree costs tens of times its size: a named error for most oversized decks
MEDIA_SUFFIXES = (".png", ".jpg", ".jpeg", ".jpe", ".gif", ".bmp", ".tif", ".tiff", ".emf", ".wmf", ".wdp", ".jxr", ".webp", ".heic",
                  ".mp3", ".mp4", ".m4a", ".m4v", ".wav", ".wma", ".wmv", ".avi", ".mov", ".mpg", ".mpeg", ".ogg",
                  ".xlsx", ".docx", ".pptx")   # what is not XML by its name; a part with any other name is counted as XML, whatever its name says
TEXT_CAP = 8 * 1024 * 1024            # what a tool's text is read up to before it is cut
CSV_DELIMITERS = ",;\t|"
SCRIPTLIKE = re.compile(r"(?is)<(script|style|noscript|template)\b.*?</\1\s*>")
PAGE_BREAK = "\f"
HIDDEN = ("0", "false")                  # the two spellings of a false xsd:boolean (a slide's show attribute)
# The tools are found on this PATH. In the container nothing sets CONVERTER_TOOL_PATH (no environment variable of the host passes), so
# it is always the default; a test on a machine that keeps pandoc elsewhere sets it.
ENV = {"PATH": os.environ.get("CONVERTER_TOOL_PATH") or "/usr/local/bin:/usr/bin:/bin", "HOME": "/tmp", "LANG": "C.UTF-8",
       "LC_ALL": "C.UTF-8"}


class Failed(Exception):
    """A conversion that cannot be made; the message goes to standard error."""


def run_tool(argv, timeout=TOOL_SECONDS):
    """One adopted tool, started with a short environment and no input; its standard output is returned as bytes (cut at
    TEXT_CAP), its failure raised as Failed with the tail of its standard error."""
    try:
        done = subprocess.run(argv, stdin=subprocess.DEVNULL, capture_output=True, timeout=timeout, env=ENV, check=False)
    except FileNotFoundError:
        raise Failed(f"{argv[0]} is not in the image") from None
    except subprocess.TimeoutExpired:
        raise Failed(f"{argv[0]} ran over {timeout} s") from None
    if done.returncode != 0:
        tail = done.stderr.decode("utf-8", "replace").strip().splitlines()[-1:] or [f"exit {done.returncode}"]
        raise Failed(f"{os.path.basename(argv[0])} failed: {tail[0][:300]}")
    return done.stdout[:TEXT_CAP]


def version_of(argv, pick=0):
    try:
        done = subprocess.run(argv, capture_output=True, timeout=30, env=ENV, check=False)
    except (OSError, subprocess.SubprocessError):
        return None
    text = (done.stdout + done.stderr).decode("utf-8", "replace").strip().splitlines()
    return text[pick].strip()[:120] if len(text) > pick else None


def utf8(data: bytes) -> str:
    """Text from bytes: UTF-8 (a BOM dropped), else Windows-1252 as a spreadsheet tool writes it; NUL characters out."""
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", "replace")
    return text.replace("\x00", "")


def check_package(path: str, bytes_max: int = 0, xml_max: int = 0) -> list:
    """The names in an Office package, refusing a zip bomb (too many entries, or a declared size that is too large; `xml_max`, when
    given, also caps the declared size of every part that is not a known kind of media or an embedded package: a part is judged by what it
    could be, not by its name, because a part saved as slide2.bin can be typed XML by an Override of [Content_Types].xml)."""
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
            if len(infos) > ZIP_ENTRIES_MAX:
                raise Failed(f"the package holds {len(infos)} entries")
            if sum(i.file_size for i in infos) > (bytes_max or ZIP_BYTES_MAX):
                raise Failed("the package declares too much data")
            if xml_max and sum(i.file_size for i in infos if not i.filename.lower().endswith(MEDIA_SUFFIXES)) > xml_max:
                raise Failed("the package declares too much XML")
            return [i.filename for i in infos]
    except zipfile.BadZipFile:
        raise Failed("the file is not a valid Office package") from None


MACRO_REFUSAL = "the package holds macros (VBA code or macro sheets): a macro-enabled Office file is not converted"
CONTENT_TYPES_CAP = 2 * 1024 * 1024     # what is read of [Content_Types].xml; a larger one is no honest package
MACRO_PART_NAMES = ("vbaproject.bin", "vbaprojectsignature.bin", "vbadata.xml")   # the last part of a part's name, lower case
MACRO_PART_FOLDERS = ("xl/macrosheets/", "xl/dialogs/")                            # Excel 4 macro sheets and dialog sheets
MACRO_TYPE_WORDS = ("macroenabled", "vbaproject", "vbadata", "macrosheet")        # in a content type, lower case
CONTENT_TYPE_VALUES = re.compile(r"""ContentType\s*=\s*(?:"([^"]*)"|'([^']*)')""", re.I)


def refuse_macros(path: str) -> None:
    """Failed with MACRO_REFUSAL when an Office package carries macros, judged by its content and never by its name (the input is
    mounted as /in/input and a sender chooses a file's name): a part named `vbaProject.bin` (also its signature and `vbaData.xml`) at
    any depth, a part under `xl/macrosheets/` or `xl/dialogs/`, or a content type in [Content_Types].xml that names a macro (any
    `macroEnabled` type, which is what .docm, .dotm, .xlsm, .xltm, .xlam, .pptm, .potm, .ppsm, .ppam and .sldm declare for their main
    part, a VBA project, a VBA data part or a macro sheet). It runs before any tool reads the file and after check_package's size
    guard. No macro is run or read: only names and the content-types list are looked at. A file that is not a package is not judged
    here (the conversion refuses it); the old binary formats (.doc, .xls, .ppt, .dot) are not a zip, are not a type the intake takes,
    and are refused by check_package if one is named."""
    names = check_package(path)
    for name in names:
        lower = name.lower()
        if lower.rsplit("/", 1)[-1] in MACRO_PART_NAMES or lower.startswith(MACRO_PART_FOLDERS):
            raise Failed(MACRO_REFUSAL)
    declared = next((n for n in names if n.lower() == "[content_types].xml"), None)
    if declared is None:
        return
    try:
        with zipfile.ZipFile(path) as z, z.open(declared) as part:
            raw = part.read(CONTENT_TYPES_CAP + 1)
    except (zipfile.BadZipFile, OSError, RuntimeError, NotImplementedError, EOFError):
        raise Failed("the package's [Content_Types].xml cannot be read") from None
    if len(raw) > CONTENT_TYPES_CAP:
        raise Failed("the package's [Content_Types].xml is too large")
    text = utf8(raw)
    for found in CONTENT_TYPE_VALUES.finditer(text):
        value = (found.group(1) or found.group(2) or "").lower()
        if any(word in value for word in MACRO_TYPE_WORDS):
            raise Failed(MACRO_REFUSAL)


def ooxml_kind(path: str) -> str:
    """docx, xlsx or pptx, from the names of the package's main parts (read in the container, never on the host). The main part decides:
    a presentation with a chart holds an embedded workbook, a document may hold a spreadsheet, so no type string is searched."""
    names = set(check_package(path))
    if "[Content_Types].xml" not in names:
        raise Failed("the package has no [Content_Types].xml")
    for main, kind in (("word/document.xml", "docx"), ("ppt/presentation.xml", "pptx"), ("xl/workbook.xml", "xlsx")):
        if main in names:
            return kind
    raise Failed("the package is not a Word document, a spreadsheet or a presentation")


def convert_pdf(path: str, work: str):
    raw = run_tool(["pdftotext", "-enc", "UTF-8", "-eol", "unix", path, "-"])
    text = utf8(raw)
    pages = text.split(PAGE_BREAK)
    if pages and not pages[-1].strip():
        pages = pages[:-1]
    out = []
    for n, page in enumerate(pages, 1):
        out.append(f"[page {n}]\n\n{page.strip()}")
    notes = []
    if not any(p.strip() for p in pages):
        notes.append("no text layer: the PDF may be scans of pages (no OCR is made)")
    return "converted.md", "\n\n".join(out).strip() + "\n", {"pdftotext": version_of(["pdftotext", "-v"])}, notes


def convert_docx(path: str, work: str):
    check_package(path)
    raw = run_tool(["pandoc", "--sandbox", "-f", "docx", "-t", "gfm", "--wrap=none", path])
    return "converted.md", utf8(raw).strip() + "\n", {"pandoc": version_of(["pandoc", "--version"])}, []


def convert_html(path: str, work: str):
    with open(path, "rb") as f:
        html = utf8(f.read(TEXT_CAP))
    html = SCRIPTLIKE.sub("", html)
    cleaned = os.path.join(work, "page.html")
    with open(cleaned, "w", encoding="utf-8") as f:
        f.write(html)
    raw = run_tool(["pandoc", "--sandbox", "-f", "html", "-t", "gfm", "--wrap=none", cleaned])
    return "converted.md", utf8(raw).strip() + "\n", {"pandoc": version_of(["pandoc", "--version"])}, \
        ["script, style and template elements were dropped before conversion"]


def delimiter_of(text: str) -> str:
    sample = text[:8192]
    try:
        return csv.Sniffer().sniff(sample, delimiters=CSV_DELIMITERS).delimiter
    except csv.Error:
        return ","


def write_csv(rows) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    size = 0
    for row in rows:
        writer.writerow(["" if v is None else v for v in row])
        size = buffer.tell()
        if size > CAP:
            break
    return buffer.getvalue()


def convert_csv(path: str, work: str):
    with open(path, "rb") as f:
        text = utf8(f.read(TEXT_CAP))
    delimiter = delimiter_of(text)
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    notes = [] if delimiter == "," else [f"the delimiter {delimiter!r} was changed to a comma"]
    return "converted.csv", write_csv(rows), {}, notes


def convert_text(path: str, work: str):
    with open(path, "rb") as f:
        text = utf8(f.read(TEXT_CAP))
    return "converted.md", text.strip() + "\n", {}, []


def cell(value):
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return value if isinstance(value, str) else str(value)


def convert_xlsx(path: str, work: str):
    check_package(path)
    try:
        import openpyxl
    except ImportError:
        raise Failed("openpyxl is not in the image") from None
    # openpyxl refuses a file whose name does not end in .xlsx (the input is mounted as /in/input); the link has a name of its own for each
    # conversion, so that a folder shared by two conversions never hands the second the first's file
    named = os.path.join(work, f"book-{secrets.token_hex(6)}.xlsx")
    os.makedirs(work, mode=0o700, exist_ok=True)
    os.symlink(path, named)
    try:
        book = openpyxl.load_workbook(named, read_only=True, data_only=True)
    except Exception as e:  # openpyxl raises its own and zipfile's errors for a file it cannot read
        raise Failed(f"the workbook cannot be read: {type(e).__name__}") from None
    notes, sheets = [], []
    for sheet in book.worksheets:
        if getattr(sheet, "sheet_state", "visible") != "visible":
            notes.append(f"the hidden sheet {sheet.title!r} was not read")
            continue
        rows, empty = [], 0
        for row in sheet.iter_rows(values_only=True):
            values = [cell(v) for v in row]
            if not any(str(v).strip() for v in values):
                empty += 1
                if empty > 1000:
                    break
                continue
            empty = 0
            rows.append(values)
            if len(rows) > 200000:
                notes.append(f"the sheet {sheet.title!r} was cut at 200000 rows")
                break
        width = max((max((i + 1 for i, v in enumerate(r) if str(v).strip()), default=0) for r in rows), default=0)
        sheets.append((sheet.title, [r[:width] for r in rows]))
    book.close()
    tools = {"openpyxl": getattr(openpyxl, "__version__", None)}
    if not sheets:
        raise Failed("the workbook holds no visible sheet with data")
    if len(sheets) == 1:
        return "converted.csv", write_csv(sheets[0][1]), tools, notes + [f"one sheet: {sheets[0][0]!r}"]
    parts = []
    for title, rows in sheets:
        parts.append(f"## Sheet: {title}\n\n```csv\n{write_csv(rows).rstrip()}\n```")
    return "converted.md", "\n\n".join(parts) + "\n", tools, notes


def forbid_network() -> None:
    """A second wall behind the container's missing network: in this process nothing can open a connection or resolve a name, so a
    converter library that tried (a URL converter, an update check) fails here before it fails there."""
    import socket

    def refuse(*args, **kwargs):
        raise OSError("the converter makes no network connection")

    socket.socket.connect = refuse
    socket.socket.connect_ex = refuse
    socket.create_connection = refuse
    socket.getaddrinfo = refuse


def table_rows(table) -> list:
    """The rows of a presentation's table as lists of cell text (a line break in a cell becomes a space, a pipe is escaped)."""
    rows = []
    for row in table.rows:
        rows.append([" ".join(c.text_frame.text.split()).replace("|", "\\|") if c.text_frame is not None else "" for c in row.cells])
    return rows


def markdown_table(rows: list) -> list:
    width = max((len(r) for r in rows), default=0)
    if not width:
        return []
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return lines


def shape_blocks(shapes, title, out: list) -> None:
    """The text of the shapes of a slide in the order they are stored: a text frame as paragraphs (a placeholder of the body as list
    items, indented by level), a table as a Markdown table, a group by its members. The title shape is the heading; pictures, charts
    and other objects hold no text that is read."""
    for shape in shapes:
        if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):    # a group of shapes
            shape_blocks(shape.shapes, title, out)
            continue
        if getattr(shape, "has_table", False) and shape.has_table:
            out.append("\n".join(markdown_table(table_rows(shape.table))))
            continue
        if (title is not None and shape.shape_id == title.shape_id) or not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
            continue
        body = shape.is_placeholder and shape.placeholder_format.type is not None and int(shape.placeholder_format.type) in (2, 7)
        lines = []
        for paragraph in shape.text_frame.paragraphs:
            text = " ".join(paragraph.text.split())
            if text:
                lines.append(("  " * paragraph.level + "- " + text) if body else text)
        if lines:
            out.append("\n".join(lines))


def convert_pptx(path: str, work: str):
    """A presentation as Markdown with python-pptx: one section per slide, headed by the slide's title (`# Slide n` when it has none),
    then its text frames, tables and, under `### Notes`, its speaker notes. A hidden slide is not read and the manifest says so.
    Nothing is fetched: python-pptx reads the one package, the process cannot open a connection (forbid_network), and a failure of the
    library is a conversion error that names the exception class and nothing else."""
    forbid_network()
    check_package(path, PPTX_BYTES_MAX, PPTX_XML_MAX)
    try:
        import pptx
        deck = pptx.Presentation(path)
        parts, notes = [], []
        for number, slide in enumerate(deck.slides, 1):
            if slide._element.get("show") in HIDDEN:    # xsd:boolean: 0 and false
                notes.append(f"the hidden slide {number} was not read")
                continue
            title = slide.shapes.title
            heading = " ".join(title.text_frame.text.split()) if title is not None and title.has_text_frame else ""
            blocks = ["# " + (heading or f"Slide {number}")]
            shape_blocks(slide.shapes, title, blocks)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                spoken = slide.notes_slide.notes_text_frame.text.strip()
                if spoken:
                    blocks.append("### Notes\n\n" + spoken)
            parts.append("\n\n".join(blocks))
            if sum(len(p) for p in parts) > CAP:
                notes.append(f"the presentation was cut after slide {number}")
                break
        text = "\n\n".join(parts).replace("\x00", "")
        version = getattr(pptx, "__version__", None)
    except Failed:
        raise
    except Exception as e:  # the library raises its own errors and its dependencies': one closed message, the class name only
        raise Failed(f"python-pptx could not convert the pptx file: {type(e).__name__}") from None
    return "converted.md", text.strip() + "\n", {"python-pptx": version}, notes


def convert(kind: str, path: str, work: str):
    if kind in ("docx", "xlsx", "pptx", "ooxml"):
        refuse_macros(path)     # for every Office package, whoever sent it and whatever the declared type says, before any tool runs
    if kind == "ooxml":
        kind = ooxml_kind(path)
    table = {"pdf": convert_pdf, "docx": convert_docx, "xlsx": convert_xlsx, "html": convert_html, "csv": convert_csv,
             "text": convert_text, "pptx": convert_pptx}
    if kind not in table:
        raise Failed(f"the type {kind!r} is not converted")
    name, text, tools, notes = table[kind](path, work)
    data = text.encode("utf-8")
    cut = len(data) > CAP
    return kind, name, data[:CAP + 1], {k: v for k, v in tools.items() if v}, notes + (["the text was cut at the cap"] if cut else [])


# --- hidden text (package P8, finding S9) ----------------------------------------------------------------------------------------
#
# A page or a PDF can carry text a person cannot see and a model reads, and the conversion hands the model all of it. This section
# finds the kinds listed below in an HTML page and in a PDF and reports a count, the kinds and short location hints, never the
# text. The host (runtime/sources.py) checks the answer against the same closed list and records it in the version's manifest and
# in its catalog.
#
#   HTML  an element's own style attribute, the simple rules of its <style> blocks (a tag, a class, an id, or a tag with a class
#         or an id) and the attributes `hidden`, `bgcolor` and `<font color>`: display:none, visibility:hidden, the hidden
#         attribute, opacity 0, font-size 0 or under 2px, a box of 1px or less with overflow hidden or a clip, a position or a
#         text-indent 500px or more off the page, and a text colour that matches its background (white on the page's default white
#         included) with a contrast ratio under 1.1. A hidden ancestor hides everything under it and is counted once.
#   PDF   poppler's pdftohtml: a fill colour that matches the white page, a font size of 2pt or less, and text in the invisible
#         render mode (the pieces that appear only when pdftohtml is told to show hidden text). A scanned PDF with a recognised-text
#         layer is flagged on purpose: the person sees the picture and the model reads that layer.
#
# What it does not find, and says so (docs/architecture/platform/outputs.md): text under a picture or a drawn shape, a colour match
# against a background that a style sheet outside the page, a selector with a combinator, a pseudo-class or an attribute, or a rule
# inside @media sets, text drawn as paths or as an image, the alt or title text of a hidden element, text a script moves, and the
# hidden text of a Word document, a spreadsheet or a presentation. A PDF's off-page text is not extracted by pdftotext, so it is
# not read and is not flagged.

HIDDEN_KINDS = ("display_none", "visibility_hidden", "hidden_attribute", "zero_size", "tiny_text", "off_page", "transparent",
                "white_on_white", "color_matched", "invisible_text")
HIDDEN_HINTS_MAX = 10
HIDDEN_OWNERS_MAX = 2000
CONTRAST_MIN = 1.1
TINY_PX, TINY_PT = 2.0, 2
OFF_PAGE_PX = 500.0
CSS_CAP = 200000
IMPLIED_END = {"p": ("p",), "li": ("li",), "dt": ("dt", "dd"), "dd": ("dt", "dd"), "tr": ("tr",), "td": ("td", "th"),
               "th": ("td", "th"), "option": ("option",)}
VOID_TAGS = frozenset("area base br col embed hr img input link meta param source track wbr".split())
UNREAD_TAGS = frozenset(("script", "style", "noscript", "template"))     # what a converter never reads, whatever it holds
HEAD_TAGS = frozenset(("html", "head", "xtitle", "meta", "link", "base", "style", "script", "noscript", "template"))
HINT_TAG_MAX, HINT_KINDS_MAX, HINT_MAX = 20, 3, 80                       # what a hint holds: the info.json that carries it stays under INFO_CAP
NAMED_COLORS = {
    "white": (255, 255, 255), "black": (0, 0, 0), "snow": (255, 250, 250), "ivory": (255, 255, 240), "whitesmoke": (245, 245, 245),
    "ghostwhite": (248, 248, 255), "floralwhite": (255, 250, 240), "azure": (240, 255, 255), "mintcream": (245, 255, 250),
    "honeydew": (240, 255, 240), "aliceblue": (240, 248, 255), "seashell": (255, 245, 238), "gainsboro": (220, 220, 220),
    "lightgray": (211, 211, 211), "lightgrey": (211, 211, 211), "silver": (192, 192, 192), "gray": (128, 128, 128),
    "grey": (128, 128, 128), "darkgray": (169, 169, 169), "darkgrey": (169, 169, 169), "red": (255, 0, 0), "blue": (0, 0, 255),
    "green": (0, 128, 0), "yellow": (255, 255, 0), "navy": (0, 0, 128),
}
LENGTH = re.compile(r"(-?\d*\.?\d+)(px|pt|em|rem|%)?")
SIMPLE_SELECTOR = re.compile(r"([a-z][a-z0-9]*|\*)?(?:\.([a-z0-9_-]+))?(?:#([a-z0-9_-]+))?")
HEX_COLOR = re.compile(r"#([0-9a-f]{3,4}|[0-9a-f]{6}|[0-9a-f]{8})")
FUNC_COLOR = re.compile(r"rgba?\(\s*([0-9.]+%?)\s*[, ]\s*([0-9.]+%?)\s*[, ]\s*([0-9.]+%?)\s*(?:[,/]\s*([0-9.]+%?)\s*)?\)")


def css_declarations(text: str) -> dict:
    """{property: value} from the text of a style attribute or of a rule's block: lower case, `!important` dropped, values cut at
    200 characters; a later declaration of a property wins."""
    out = {}
    for part in text[:4000].split(";"):
        if ":" not in part:
            continue
        name, value = part.split(":", 1)
        name = name.strip().lower()
        if re.fullmatch(r"-?[a-z][a-z-]{0,40}", name):
            out[name] = re.sub(r"\s*!important\s*$", "", value.strip().lower())[:200]
    return out


def css_rules(css: str) -> list:
    """The simple rules of the style blocks, [(rank, tag, class, id, declarations)] in the order a cascade applies them (an id over
    a class over a tag, later over earlier). Rules inside an at-rule (@media, @supports) and selectors with a combinator, a
    pseudo-class or an attribute are left out."""
    css = re.sub(r"/\*.*?\*/", "", css[:CSS_CAP], flags=re.S)
    css = re.sub(r"@[a-z-]+[^{};]*\{(?:[^{}]*\{[^{}]*\})*[^{}]*\}", "", css)
    rules = []
    for block in re.finditer(r"([^{}]+)\{([^{}]*)\}", css):
        declarations = css_declarations(block.group(2))
        for selector in block.group(1).split(","):
            found = SIMPLE_SELECTOR.fullmatch(selector.strip().lower())
            if found and selector.strip():
                tag, klass, ident = found.groups()
                rank = (3 if ident else 0) + (2 if klass else 0) + (1 if tag and tag != "*" else 0)
                rules.append((rank, None if tag == "*" else tag, klass, ident, declarations))
    return sorted(rules, key=lambda r: r[0])


def css_length(value, em: float = 16.0):
    """A CSS length in pixels, or None when it is not a plain one (`auto`, a calc, a viewport unit, a percentage)."""
    found = LENGTH.fullmatch((value or "").strip())
    if not found:
        return None
    number, unit = float(found.group(1)), found.group(2)
    if unit in (None, "px"):
        return number
    if unit == "pt":
        return number * 4.0 / 3.0
    if unit in ("em", "rem"):
        return number * em
    return None


def css_color(value):
    """(r, g, b, alpha) of a colour written as a name, #rgb, #rgba, #rrggbb, #rrggbbaa, rgb() or rgba(); None for anything else
    (inherit, currentcolor, hsl, a variable)."""
    text = (value or "").strip().lower()
    if text == "transparent":
        return (0, 0, 0, 0.0)
    if text in NAMED_COLORS:
        return NAMED_COLORS[text] + (1.0,)
    hexed = HEX_COLOR.fullmatch(text)
    if hexed:
        digits = hexed.group(1)
        if len(digits) in (3, 4):
            digits = "".join(c * 2 for c in digits)
        alpha = int(digits[6:8], 16) / 255.0 if len(digits) == 8 else 1.0
        return (int(digits[0:2], 16), int(digits[2:4], 16), int(digits[4:6], 16), alpha)
    func = FUNC_COLOR.fullmatch(text)
    if func:
        def channel(part):
            return min(255, max(0, round(float(part[:-1]) * 2.55 if part.endswith("%") else float(part))))
        alpha = 1.0
        if func.group(4):
            alpha = float(func.group(4)[:-1]) / 100.0 if func.group(4).endswith("%") else float(func.group(4))
        return (channel(func.group(1)), channel(func.group(2)), channel(func.group(3)), alpha)
    return None


def luminance(rgb) -> float:
    def part(c):
        c = c / 255.0
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    return 0.2126 * part(rgb[0]) + 0.7152 * part(rgb[1]) + 0.0722 * part(rgb[2])


def contrast(a, b) -> float:
    """The WCAG contrast ratio of two colours: 1.0 (the same) to 21.0."""
    high, low = sorted((luminance(a), luminance(b)), reverse=True)
    return (high + 0.05) / (low + 0.05)


def color_reason(color, background):
    """The kind of hidden text a text colour on a background is, or None. `color` is None for the default (black); `background` is
    None for the page's default white, or "unknown" for one that cannot be worked out (a picture, a gradient): then nothing is said."""
    if background == "unknown":
        return None
    if color is not None and color[3] <= 0.05:
        return "transparent"
    ink = color[:3] if color is not None else (0, 0, 0)
    paper = background[:3] if background is not None and background[3] > 0.05 else (255, 255, 255)
    if contrast(ink, paper) >= CONTRAST_MIN:
        return None
    return "white_on_white" if background is None else "color_matched"


def hint_of(owner: dict, number: int) -> str:
    """The location hint of a hidden element: its kinds (the first HINT_KINDS_MAX), its tag name reduced to letters, digits and hyphens
    and cut at HINT_TAG_MAX, and its number in the document; at most HINT_MAX characters whatever the page names its tags."""
    tag = re.sub(r"[^a-z0-9-]", "", owner["tag"].lower())[:HINT_TAG_MAX] or "x"
    kinds = "+".join([k for k in HIDDEN_KINDS if k in owner["kinds"]][:HINT_KINDS_MAX])
    return f"{kinds} in {tag} element {number}"[:HINT_MAX]


class HiddenScan:
    """What an HTML page hides: feed(html), then result(). A small walk over html.parser's events that keeps, for each open element,
    the reasons its text is hidden (those of the elements around it and its own) and the colour, size and background it resolves
    to; text found under an element that has a reason is counted once, for the outermost element that has one."""

    def __init__(self, rules):
        import html.parser

        outer = self

        class Parser(html.parser.HTMLParser):
            def handle_starttag(self, tag, attrs):
                outer.start(tag, attrs)

            def handle_startendtag(self, tag, attrs):
                outer.start(tag, attrs)
                outer.end(tag)

            def handle_endtag(self, tag):
                outer.end(tag)

            def handle_data(self, data):
                outer.data(data)

        self.rules = rules
        self.parser = Parser(convert_charrefs=True)
        self.stack = []
        self.number = 0
        self.unread = 0
        self.in_head = False
        self.body_seen = False
        self.title_open = False
        self.title_hidden = None
        self.capped = False
        self.owners = {}

    def feed(self, html: str) -> None:
        # Newer versions of html.parser read the content of a <title> as plain text up to </title> (as the HTML standard does), so a title
        # that is never closed would swallow the page. A converter reads the tags either way, so the title is renamed before the parse.
        html = re.sub(r"(?i)<(/?)title(?=[\s>/])", r"<\1xtitle", html)
        self.parser.feed(html)
        self.parser.close()

    def own_style(self, tag, attrs) -> dict:
        classes = set((attrs.get("class") or "").lower().split())
        ident = (attrs.get("id") or "").lower()
        found = {}
        for _rank, rtag, rclass, rid, declarations in self.rules:
            if (rtag is None or rtag == tag) and (rclass is None or rclass in classes) and (rid is None or rid == ident):
                found.update(declarations)
        found.update(css_declarations(attrs.get("style") or ""))
        return found

    def start(self, tag, pairs) -> None:
        self.number += 1
        if tag == "xtitle":
            # A title is never shown (the browser's own style sheet hides it), wherever it stands. Its text outside the head, and the
            # rest of a page after a title that was never closed (a browser reads it all as the title), is hidden text; a converter
            # that reads the tags reads it as the page's text.
            self.title_open = True
            self.title_owner = {"n": self.number, "tag": "title"}
            if self.body_seen:
                self.title_hidden = self.title_owner
        if tag in ("head", "xtitle"):
            self.in_head = True
        elif self.in_head and tag not in HEAD_TAGS:
            self.in_head = False                  # a head or a title that was never closed ends where the body's first tag begins
            self.body_seen = True
            if self.title_open:
                self.title_hidden = self.title_owner
        elif tag not in HEAD_TAGS:
            self.body_seen = True
        if tag in UNREAD_TAGS:
            self.unread += 1
        if tag in VOID_TAGS:
            return
        if self.stack and self.stack[-1]["tag"] in IMPLIED_END.get(tag, ()):
            self.stack.pop()                      # <p>a<p>b: the first paragraph ended where the second began
        attrs = {k.lower(): (v or "") for k, v in pairs[:60] if isinstance(k, str)}
        style = self.own_style(tag, attrs)
        root = {"hide": frozenset(), "color": None, "font": None, "background": None, "visibility": False}
        parent = self.stack[-1] if self.stack else root
        hide = set(parent["hide"])
        if "hidden" in attrs:
            hide.add("hidden_attribute")
        if (style.get("display") or "").split(" ")[0] == "none":
            hide.add("display_none")
        visibility = parent["visibility"]
        if style.get("visibility") in ("hidden", "collapse"):
            visibility = True
        elif style.get("visibility") == "visible":
            visibility = False
        try:
            if "opacity" in style and float(style["opacity"]) <= 0.01:
                hide.add("transparent")
        except ValueError:
            pass
        small = [css_length(style.get(p)) for p in ("width", "height", "max-width", "max-height") if p in style]
        clipped = style.get("overflow") in ("hidden", "clip") or "clip" in style or "clip-path" in style
        if clipped and any(v is not None and v <= 1.0 for v in small):
            hide.add("zero_size")
        edges = [css_length(style.get(p)) for p in ("left", "right", "top", "bottom") if p in style]
        indent = css_length(style.get("text-indent"))
        if (style.get("position") in ("absolute", "fixed") and any(v is not None and v <= -OFF_PAGE_PX for v in edges)) \
                or (indent is not None and indent <= -OFF_PAGE_PX):
            hide.add("off_page")
        font = parent["font"]
        if "font-size" in style:
            size = style["font-size"]
            if size.endswith("%"):
                try:
                    font = (font if font is not None else 16.0) * float(size[:-1]) / 100.0
                except ValueError:
                    pass
            else:
                length = css_length(size, em=font if font is not None else 16.0)
                font = length if length is not None else font
        color = parent["color"]
        wanted = style.get("color") or (attrs.get("color") if tag == "font" else None)
        if wanted:
            color = css_color(wanted) or color
        background = parent["background"]
        painted = style.get("background-color") or style.get("background") or attrs.get("bgcolor")
        if painted:
            first = painted.split(" ")[0]
            parsed = None if ("url(" in painted or "gradient" in painted) else (css_color(painted) or css_color(first))
            background = parsed if parsed is not None else "unknown"
        reasons = set(hide)
        if visibility:
            reasons.add("visibility_hidden")
        if font is not None and font < TINY_PX:
            reasons.add("zero_size" if font <= 0 else "tiny_text")
        reason = color_reason(color, background) if (color is not None or background not in (None, "unknown")) else None
        if reason:
            reasons.add(reason)
        self.stack.append({"tag": tag, "n": self.number, "hide": frozenset(hide), "color": color, "font": font,
                           "background": background, "visibility": visibility, "reasons": frozenset(reasons)})

    def end(self, tag) -> None:
        if tag in UNREAD_TAGS and self.unread:
            self.unread -= 1
        if tag == "xtitle":
            self.title_open = False
            self.title_hidden = None
        if tag in ("head", "xtitle"):
            self.in_head = False
        for depth in range(len(self.stack) - 1, max(-1, len(self.stack) - 40), -1):
            if self.stack[depth]["tag"] == tag:
                del self.stack[depth:]
                return

    def note(self, owner: dict, reasons) -> None:
        if owner["n"] not in self.owners and len(self.owners) >= HIDDEN_OWNERS_MAX:
            self.capped = True
            return
        self.owners.setdefault(owner["n"], {"tag": owner["tag"], "kinds": set()})["kinds"] |= set(reasons)

    def data(self, text) -> None:
        if self.unread or not text.strip():
            return
        if self.title_hidden is not None:
            self.note(self.title_hidden, ("display_none",))
            return
        if self.in_head or not self.stack or not self.stack[-1]["reasons"]:
            return
        self.note(next(e for e in self.stack if e["reasons"]), self.stack[-1]["reasons"])

    def result(self) -> dict:
        kinds = {}
        for owner in self.owners.values():
            for kind in owner["kinds"]:
                kinds[kind] = kinds.get(kind, 0) + 1
        hints = [hint_of(owner, n) for n, owner in sorted(self.owners.items())[:HIDDEN_HINTS_MAX]]
        record = {"examined": True, "count": len(self.owners), "kinds": {k: kinds[k] for k in HIDDEN_KINDS if k in kinds}, "hints": hints}
        if self.capped:
            record["capped"] = True               # the count stops at HIDDEN_OWNERS_MAX: there are at least that many
        return record


def hidden_in_html(html: str) -> dict:
    """The hidden-text record of an HTML page: {"examined": True, "count", "kinds": {kind: n}, "hints": [...]}."""
    html = html[:TEXT_CAP]
    scan = HiddenScan(css_rules("\n".join(re.findall(r"(?is)<style\b[^>]*>(.*?)</style\s*>", html))))
    scan.feed(html)
    return scan.result()


PDF_PAGE = re.compile(r'<page number="(\d+)"[^>]*? height="(\d+)" width="(\d+)"')
PDF_FONT = re.compile(r'<fontspec id="(\d+)" size="(-?\d+)" family="[^"]*" color="(#[0-9a-fA-F]{6})"')
PDF_TEXT = re.compile(r'<text top="(-?\d+)" left="(-?\d+)" width="(-?\d+)" height="(-?\d+)" font="(\d+)">(.*?)</text>')


def pdf_texts(xml: str) -> list:
    """[(page, top, left, text, size in pt, colour)] of the pieces of text in pdftohtml's XML that lie on their page. A piece off
    the page is not read by pdftotext, so it is left out."""
    fonts, out = {}, []
    page = height = width = 0
    for line in xml.splitlines():
        found = PDF_PAGE.search(line)
        if found:
            page, height, width = int(found.group(1)), int(found.group(2)), int(found.group(3))
            continue
        found = PDF_FONT.search(line)
        if found:
            fonts[found.group(1)] = (int(found.group(2)), found.group(3).lower())
            continue
        found = PDF_TEXT.search(line)
        if found and found.group(5) in fonts:
            top, left, w, h = (int(found.group(i)) for i in (1, 2, 3, 4))
            text = re.sub(r"<[^>]*>", "", found.group(6)).strip()
            if text and not (left + w <= 0 or left >= width or top + h <= 0 or top >= height):
                out.append((page, top, left, text, fonts[found.group(5)][0], fonts[found.group(5)][1]))
    return out


def hidden_in_pdf(shown_xml: str, all_xml: str) -> dict:
    """The hidden-text record of a PDF from two runs of pdftohtml -xml: `shown_xml` as it comes and `all_xml` with -hidden. A piece
    of text is flagged when its fill colour matches the white page or its font is 2pt or less; a piece that only the second run
    has is in the invisible render mode."""
    visible = pdf_texts(shown_xml)
    flagged = []                                  # [(page, {kinds})], one for each piece
    for page, _top, _left, _text, size, color in visible:
        reasons = set()
        if size <= TINY_PT:
            reasons.add("zero_size" if size <= 0 else "tiny_text")
        if contrast((int(color[1:3], 16), int(color[3:5], 16), int(color[5:7], 16)), (255, 255, 255)) < CONTRAST_MIN:
            reasons.add("white_on_white")
        if reasons:
            flagged.append((page, reasons))
    seen = {}
    for page, top, left, text, _size, _color in visible:
        seen[(page, top, left, text)] = seen.get((page, top, left, text), 0) + 1
    for page, top, left, text, _size, _color in pdf_texts(all_xml):
        if seen.get((page, top, left, text), 0) > 0:
            seen[(page, top, left, text)] -= 1
        else:
            flagged.append((page, {"invisible_text"}))
    kinds, pages = {}, {}
    for page, reasons in flagged:
        for kind in reasons:
            kinds[kind] = kinds.get(kind, 0) + 1
            pages[(page, kind)] = pages.get((page, kind), 0) + 1
    hints = [f"{kind} on page {page} count {n}" for (page, kind), n in sorted(pages.items(), key=lambda i: (i[0][0], HIDDEN_KINDS.index(i[0][1])))]
    return {"examined": True, "count": len(flagged), "kinds": {k: kinds[k] for k in HIDDEN_KINDS if k in kinds},
            "hints": hints[:HIDDEN_HINTS_MAX]}


def hidden_text(kind: str, path: str):
    """The hidden-text record of a converted file: for an HTML page and a PDF the record; for any other kind None (not examined). A
    check that cannot be made is {"examined": False, "why": "tool_failed"}: the conversion is not failed for it, and the host says so."""
    try:
        if kind == "html":
            with open(path, "rb") as f:
                return hidden_in_html(utf8(f.read(TEXT_CAP)))
        if kind == "pdf":
            base = ["pdftohtml", "-xml", "-zoom", "1", "-i", "-stdout", "-enc", "UTF-8"]
            shown = utf8(run_tool(base + [path]))
            everything = utf8(run_tool(base[:2] + ["-hidden"] + base[2:] + [path]))
            return hidden_in_pdf(shown, everything)
    except (Failed, OSError, RecursionError):
        return {"examined": False, "why": "tool_failed"}
    return None


def bounded_info(info: dict) -> bytes:
    """The info.json text, always whole JSON and at most INFO_CAP bytes: when it is longer, the parts that can go go in order (the hidden
    text's hints, the notes, the tool versions) and never the text cut in the middle, which the host could not parse. The hidden-text
    record keeps its count and its kinds, which are bounded by their number."""
    info = dict(info)
    for drop in ("hints", "notes", "tool_versions", None):
        data = json.dumps(info, sort_keys=True).encode("utf-8")
        if len(data) <= INFO_CAP:
            return data
        if drop == "hints" and isinstance(info.get("hidden_text"), dict):
            info["hidden_text"] = dict(info["hidden_text"], hints=[])
        elif drop in ("notes", "tool_versions"):
            info[drop] = [] if drop == "notes" else {}
    raise Failed("the account of the conversion is too long")


def frames(files) -> bytes:
    out = [b"OHF1\n"]
    for name, data in files:
        out.append(f"FILE {name} {len(data)}\n".encode("ascii") + data)
    out.append(b"END\n")
    return b"".join(out)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--type", required=True)
    parser.add_argument("--input", required=True)
    parser.add_argument("--work", default="/scratch/work")
    args = parser.parse_args(argv)
    if not args.input.startswith("/"):    # a path in the container; an address (http://, file:, data:) is never an input
        print("error: the input is a file path, never an address", file=sys.stderr)
        return 3
    os.makedirs(args.work, mode=0o700, exist_ok=True)
    try:
        kind, name, data, tools, notes = convert(args.type, args.input, args.work)
    except Failed as e:
        print(f"error: {e}", file=sys.stderr)
        return 3
    except OSError as e:
        print(f"error: {type(e).__name__}: {e.strerror}", file=sys.stderr)
        return 1
    info = bounded_info({"converter": "document", "kind": kind, "tool_versions": tools, "notes": notes,
                         "hidden_text": hidden_text(kind, args.input)})
    sys.stdout.buffer.write(frames([(name, data), ("info.json", info)]))
    sys.stdout.buffer.flush()
    shutil.rmtree(args.work, ignore_errors=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
